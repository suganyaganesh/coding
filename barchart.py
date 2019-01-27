from openpyxl import Workbook
from openpyxl.chart import BarChart,Series,Reference
import os

book1=Workbook(write_only=True)
sheet=book1.create_sheet()


rows=[('Numbers','Batch1','Batch2'),(2,10,30),(3,40,60),(4,50,70),(5,20,10),(6,10,40),(7,50,30)]

for row in rows:
 sheet.append(row)
 
chart1=BarChart()
chart1.type="col"
chart1.style=10
chart1.title="Bar Chart"
chart1.y_axis.title='Test number'
chart1.x_axis.title='Sample Length(mm)'

data=Reference(sheet,min_col=2,min_row=1,max_row=7,max_col=3)
cats=Reference(sheet,min_col=1,min_row=2,max_row=7)
chart1.add_data(data,titles_from_data=True)
chart1.set_categories(cats)
chart1.shape=4
sheet.add_chart(chart1,"A10")

book1.save("bar.xlsx")

os.system('start bar.xlsx')